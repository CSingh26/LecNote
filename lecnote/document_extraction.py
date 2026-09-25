"""Bounded, read-only extraction of text and Office Open XML attachments.

Office packages are validated in memory before third-party parsers see them.
No package member is written to disk; embedded objects and external links are
never opened. Formula strings are text, never calculations.
"""

import codecs
import io
import re
import struct
import unicodedata
import zipfile
from collections.abc import Iterable
from pathlib import Path, PurePosixPath

MAX_TEXT_CHARS = 100_000
MAX_INPUT_BYTES = 30 * 1024 * 1024
MAX_ZIP_MEMBERS = 2048
MAX_MEMBER_BYTES = 8 * 1024 * 1024
MAX_EXPANDED_BYTES = 32 * 1024 * 1024
MAX_XML_ELEMENTS = 250_000
MAX_XML_DEPTH = 64
MAX_SHEET_ROWS = 10_000
MAX_SHEET_COLUMNS = 256
MAX_SHEET_CELLS = 100_000
MAX_TABLE_SPAN = 256

_WORD_NAMESPACE = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_SHEET_NAMESPACE = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

_OFFICE_PARTS = {
    ".docx": "word/document.xml",
    ".pptx": "ppt/presentation.xml",
    ".xlsx": "xl/workbook.xml",
}
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_BINARY_MAGIC = (
    b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08", b"\x1f\x8b", b"BZh", b"\xfd7zXZ\x00",
    b"7z\xbc\xaf\x27\x1c", b"Rar!", b"!<arch>\n", b"\x7fELF", b"MZ", _OLE_MAGIC,
    b"\xfe\xed\xfa\xce", b"\xce\xfa\xed\xfe", b"\xfe\xed\xfa\xcf", b"\xcf\xfa\xed\xfe",
    b"\xca\xfe\xba\xbe", b"\x00asm",
)


def bounded_text(parts: Iterable[str], separator: str = "\n") -> str:
    """Consume text lazily, stopping at the shared character budget."""
    chunks = []
    remaining = MAX_TEXT_CHARS
    for part in parts:
        if not part:
            continue
        if chunks:
            prefix = separator[:remaining]
            chunks.append(prefix)
            remaining -= len(prefix)
        chunk = part[:remaining]
        chunks.append(chunk)
        remaining -= len(chunk)
        if not remaining:
            break
    return "".join(chunks)


def extract_document(path: Path) -> str:
    """Extract supported Office documents or strictly decoded readable text."""
    try:
        with path.open("rb") as source:
            data = source.read(MAX_INPUT_BYTES + 1)
        if len(data) > MAX_INPUT_BYTES:
            raise RuntimeError("Attachment exceeds the 30 MiB extraction input limit.")
        if path.suffix.lower() in _OFFICE_PARTS:
            text = _extract_office(data, path.suffix.lower())
        else:
            text = _extract_text(data)
    except OSError as exc:
        raise RuntimeError("Cannot read attachment. Check local file permissions.") from exc
    if not text.strip():
        raise RuntimeError("Attachment is empty or has no readable text.")
    return text


def _extract_text(data: bytes) -> str:
    if data.startswith(_BINARY_MAGIC) or data[257:262] == b"ustar":
        raise RuntimeError("Unsupported binary, archive, or executable content cannot be extracted as text.")
    encoding = "utf-8-sig"
    if data.startswith((codecs.BOM_UTF32_LE, codecs.BOM_UTF32_BE)):
        encoding = "utf-32"
    elif data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        encoding = "utf-16"
    decoder = codecs.getincrementaldecoder(encoding)(errors="strict")
    chunks = []
    remaining = MAX_TEXT_CHARS
    try:
        for offset in range(0, len(data), 65_536):
            text = decoder.decode(data[offset:offset + 65_536], final=offset + 65_536 >= len(data))
            if any(unicodedata.category(char) == "Cc" and char not in "\t\n\r\f" for char in text):
                raise RuntimeError("Binary control bytes found; attachment is not readable text.")
            if remaining:
                chunks.append(text[:remaining])
                remaining -= len(chunks[-1])
    except UnicodeError as exc:
        raise RuntimeError("Not readable UTF-8 text or BOM-marked UTF-16/UTF-32 text.") from exc
    return "".join(chunks)


def _check_zip_directory(data: bytes) -> None:
    # Bound the central directory before ZipFile allocates a ZipInfo per entry.
    offset = data.rfind(b"PK\x05\x06", max(0, len(data) - 65_557))
    if offset < 0 or offset + 22 > len(data):
        raise RuntimeError("Unsupported or invalid Office ZIP package.")
    disk, directory_disk, disk_count, count, size, start, comment = struct.unpack_from("<4H2LH", data, offset + 4)
    if offset + 22 + comment != len(data) or disk or directory_disk or disk_count != count:
        raise RuntimeError("Unsupported or invalid multipart Office ZIP package.")
    if count > MAX_ZIP_MEMBERS or size > MAX_ZIP_MEMBERS * 1024:
        raise RuntimeError("Office ZIP directory exceeds the member count or size limit.")
    if start + size != offset or data[max(0, offset - 20):offset - 16] == b"PK\x06\x07":
        raise RuntimeError("Unsupported or invalid Office ZIP layout (including ZIP64).")


def _validate_xml(data: bytes, *, sheet: bool, elements: int) -> int:
    from defusedxml.ElementTree import iterparse

    depth = 0
    for event, element in iterparse(io.BytesIO(data), events=("start", "end"), forbid_dtd=True):
        if event == "start":
            elements += 1
            depth += 1
            if elements > MAX_XML_ELEMENTS or depth > MAX_XML_DEPTH:
                raise RuntimeError("Office XML exceeds the element count or depth limit.")
            if element.tag in {_WORD_NAMESPACE + name for name in ("gridSpan", "gridBefore", "gridAfter")}:
                span = int(element.get(_WORD_NAMESPACE + "val", "0"))
                if not 0 <= span <= MAX_TABLE_SPAN or (element.tag.endswith("}gridSpan") and not span):
                    raise RuntimeError("Word table dimensions exceed the span limit for extraction.")
            if sheet and element.tag.startswith(_SHEET_NAMESPACE):
                name = element.tag.rsplit("}", 1)[-1]
                if name == "row" and int(element.get("r", "0")) > MAX_SHEET_ROWS:
                    raise RuntimeError("Spreadsheet exceeds the row limit for extraction.")
                if name == "c":
                    reference = element.get("r", "")
                    match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]*)", reference)
                    if not match:
                        raise RuntimeError("Invalid spreadsheet cell reference.")
                    column = 0
                    for char in match[1]:
                        column = column * 26 + ord(char) - ord("A") + 1
                    if column > MAX_SHEET_COLUMNS or int(match[2]) > MAX_SHEET_ROWS:
                        raise RuntimeError("Spreadsheet exceeds the row or column limit for extraction.")
        else:
            depth -= 1
            element.clear()
    return elements


def _validate_office(source: io.BytesIO, suffix: str) -> None:
    from defusedxml.common import DefusedXmlException
    from defusedxml.ElementTree import fromstring

    with zipfile.ZipFile(source) as package:
        members = package.infolist()
        if len(members) > MAX_ZIP_MEMBERS:
            raise RuntimeError("Office ZIP has too many members for extraction.")
        names = {item.filename for item in members}
        if not {"[Content_Types].xml", "_rels/.rels", _OFFICE_PARTS[suffix]} <= names:
            raise RuntimeError("Unsupported or invalid Office document package.")
        if len(names) != len(members):
            raise RuntimeError("Invalid Office ZIP contains duplicate members.")
        total = 0
        for item in members:
            name = PurePosixPath(item.filename)
            if name.is_absolute() or ".." in name.parts or "\\" in item.filename or "\x00" in item.filename:
                raise RuntimeError("Invalid Office ZIP member path.")
            if item.flag_bits & 1:
                raise RuntimeError("Office ZIP is encrypted. Export an unlocked copy to extract text.")
            if item.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
                raise RuntimeError("Unsupported Office ZIP compression.")
            total += item.file_size
            if item.file_size > MAX_MEMBER_BYTES or total > MAX_EXPANDED_BYTES:
                raise RuntimeError("Office ZIP exceeds the expanded member or total byte limit.")

        content_types_data = package.read("[Content_Types].xml")
        elements = _validate_xml(content_types_data, sheet=False, elements=0)
        content_types = fromstring(content_types_data, forbid_dtd=True)
        defaults = {node.get("Extension"): node.get("ContentType", "") for node in content_types
                    if node.tag.endswith("}Default")}
        overrides = {node.get("PartName", "").lstrip("/"): node.get("ContentType", "")
                     for node in content_types if node.tag.endswith("}Override")}
        for item in members:
            if item.filename == "[Content_Types].xml":
                continue
            with package.open(item) as member:
                data = member.read(MAX_MEMBER_BYTES + 1)
            if len(data) != item.file_size or len(data) > MAX_MEMBER_BYTES:
                raise RuntimeError("Office ZIP member exceeds its size limit or is damaged.")
            # Content types cover XML under unusual extensions. Opaque printer
            # settings/media can contain XML fragments but are never parsed.
            declaration_bytes = data.replace(b"\x00", b"").lower()
            content_type = overrides.get(item.filename, defaults.get(item.filename.rsplit(".", 1)[-1], ""))
            is_xml = (
                content_type.endswith(("+xml", "/xml"))
                or item.filename.lower().endswith((".xml", ".rels"))
                or declaration_bytes.lstrip(b"\xef\xbb\xbf\xff\xfe \t\r\n").startswith(b"<")
            )
            if is_xml:
                if b"<!doctype" in declaration_bytes or b"<!entity" in declaration_bytes:
                    raise RuntimeError("Unsafe Office XML DTD/entity declarations are unsupported.")
                try:
                    elements = _validate_xml(
                        data, sheet=suffix == ".xlsx",
                        elements=elements,
                    )
                except DefusedXmlException as exc:
                    raise RuntimeError("Unsafe Office XML entities or DTD are unsupported.") from exc


def _extract_office(data: bytes, suffix: str) -> str:
    if data.startswith(_OLE_MAGIC):
        raise RuntimeError("Encrypted or legacy Office document. Export an unlocked DOCX, PPTX, or XLSX copy.")
    try:
        _check_zip_directory(data)
        with io.BytesIO(data) as source:
            _validate_office(source, suffix)
            source.seek(0)
            if suffix == ".docx":
                return bounded_text(_docx_parts(source))
            if suffix == ".pptx":
                return bounded_text(_pptx_parts(source))
            return _xlsx_text(source)
    except ImportError as exc:
        raise RuntimeError("Office extraction requires LecNote's document parser dependencies. Reinstall them.") from exc
    except RuntimeError:
        raise
    except Exception as exc:
        raise RuntimeError("Unsupported or invalid Office document. It may be damaged; export a new copy.") from exc


def _docx_parts(source: io.BytesIO) -> Iterable[str]:
    from docx import Document
    from docx.table import Table, _Cell

    def table_blocks(table):
        # row.cells expands gridSpan and recursively resolves vMerge. Iterate
        # physical cells instead, so hostile dimensions cannot multiply work.
        for row in table._tbl.iterchildren(_WORD_NAMESPACE + "tr"):
            for cell in row.iterchildren(_WORD_NAMESPACE + "tc"):
                yield from _Cell(cell, table).iter_inner_content()

    document = Document(source)
    stack = [iter(document.iter_inner_content())]
    while stack:
        item = next(stack[-1], None)
        if item is None:
            stack.pop()
        elif isinstance(item, Table):
            stack.append(iter(table_blocks(item)))
        else:
            yield item.text


def _pptx_parts(source: io.BytesIO) -> Iterable[str]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for slide in Presentation(source).slides:
        stack = [iter(slide.shapes)]
        while stack:
            shape = next(stack[-1], None)
            if shape is None:
                stack.pop()
            elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                stack.append(iter(shape.shapes))
            elif shape.has_text_frame:
                yield shape.text_frame.text
            elif shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        yield cell.text
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            yield slide.notes_slide.notes_text_frame.text


def _xlsx_text(source: io.BytesIO) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False, keep_vba=False)
    try:
        def parts():
            visited = 0
            for sheet in workbook.worksheets:
                # Ignore untrusted dimension metadata; XML cell coordinates were checked above.
                sheet.reset_dimensions()
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if row_number > MAX_SHEET_ROWS:
                        raise RuntimeError("Spreadsheet exceeds the row traversal limit for extraction.")
                    visited += len(row)
                    if visited > MAX_SHEET_CELLS:
                        raise RuntimeError("Spreadsheet exceeds the cell traversal limit for extraction.")
                    if any(value is not None for value in row):
                        yield "\t".join("" if value is None else str(value) for value in row)
        return bounded_text(parts())
    finally:
        workbook.close()
