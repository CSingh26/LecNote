"""Real document fixtures and adversarial inputs for local extraction."""

import socket
import struct
import subprocess
import zipfile

import pytest

from lecnote.enrichment import extract_context


def office_file(path, *, text="Lecture material"):
    if path.suffix.lower() == ".docx":
        from docx import Document

        document = Document()
        document.add_paragraph(text)
        document.add_table(rows=1, cols=2).rows[0].cells[1].text = "Table content" if text else ""
        document.add_paragraph("After table" if text else "")
        document.save(path)
    elif path.suffix.lower() == ".pptx":
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = text
        slide.shapes.add_table(1, 2, 0, Inches(2), Inches(4), Inches(1)).table.cell(0, 1).text = (
            "Table content" if text else ""
        )
        presentation.save(path)
    else:
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append([text, 42 if text else None])
        second = workbook.create_sheet("Second sheet")
        second.append(["Table content" if text else None])
        workbook.save(path)
        workbook.close()
    return path


def rewrite_zip(path, replacements=None, additions=()):
    replacements = replacements or {}
    with zipfile.ZipFile(path) as source:
        members = [(item.filename, source.read(item)) for item in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in members:
            target.writestr(name, replacements.get(name, data))
        for name, data in additions:
            target.writestr(name, data)


@pytest.mark.parametrize("name", ["example.py", "source.rs", "custom.whatever", "Makefile", ".gitignore", "data.csv"])
def test_readable_content_ignores_extension_and_is_not_executed(tmp_path, name):
    source = tmp_path / name
    text = 'import os\nos.remove("source")\n# Caf\u00e9\nname,value\n"one,two",3\n'
    source.write_text(text, encoding="utf-8")
    before = source.read_bytes()
    assert extract_context(source) == text
    assert source.read_bytes() == before


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "utf-32"])
def test_bom_marked_unicode_text(tmp_path, encoding):
    source = tmp_path / "text"
    source.write_bytes("Lecture \u03b1\nSecond line".encode(encoding))
    assert extract_context(source) == "Lecture \u03b1\nSecond line"


@pytest.mark.parametrize("suffix", [".docx", ".PPTX", ".xlsx"])
def test_real_office_documents_extract_content_in_order_and_preserve_source(tmp_path, suffix):
    source = office_file(tmp_path / ("lecture" + suffix))
    before = source.read_bytes()
    text = extract_context(source)
    assert text.index("Lecture material") < text.index("Table content")
    if suffix == ".docx":
        assert text.index("Table content") < text.index("After table")
    assert source.read_bytes() == before


@pytest.mark.parametrize("suffix", [".docx", ".pptx", ".xlsx"])
def test_empty_office_documents_are_download_only(tmp_path, suffix):
    source = office_file(tmp_path / ("empty" + suffix), text="")
    with pytest.raises(RuntimeError, match="(?i)(empty|no readable|no extractable).*download-only"):
        extract_context(source)


@pytest.mark.parametrize("name,content", [
    ("binary.unknown", b"\x00\x01\x02binary"),
    ("fake.txt", b"\xff\x80\xfe"),
    ("executable", b"\x7fELF" + b"\0" * 100),
    ("program.exe", b"MZ" + b"\0" * 100),
    ("empty", b""),
    ("blank", b"  \n\t"),
    ("broken.docx", b"This is not an Office document"),
    ("broken.pptx", b"PK\x03\x04broken"),
    ("broken.xlsx", b"PK\x03\x04broken"),
    ("encrypted.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100),
    ("archive", b"!<arch>\n" + b" " * 100),
])
def test_unreadable_inputs_are_preserved_with_download_only_error(tmp_path, name, content):
    source = tmp_path / name
    source.write_bytes(content)
    with pytest.raises(RuntimeError, match="(?i)stored.*download-only"):
        extract_context(source)
    assert source.read_bytes() == content


@pytest.mark.parametrize("name", ["archive.zip", "archive", "archive.txt", "archive.docx"])
def test_arbitrary_archives_are_not_unpacked(tmp_path, name):
    source = tmp_path / name
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("../escape.txt", "Do not extract me")
    before = source.read_bytes()
    with pytest.raises(RuntimeError, match="download-only"):
        extract_context(source)
    assert source.read_bytes() == before
    assert list(tmp_path.iterdir()) == [source]


def test_text_output_is_bounded_but_invalid_tail_is_still_rejected(tmp_path):
    source = tmp_path / "large.code"
    source.write_text("\u03b1" * 120_000, encoding="utf-8")
    assert extract_context(source) == "\u03b1" * 100_000
    with source.open("ab") as output:
        output.write(b"\x00binary tail")
    with pytest.raises(RuntimeError, match="download-only"):
        extract_context(source)


@pytest.mark.parametrize("suffix", [".docx", ".pptx", ".xlsx"])
def test_office_output_is_bounded(tmp_path, suffix):
    source = office_file(tmp_path / ("large" + suffix), text="word " * 28_000)
    if suffix == ".xlsx":
        from openpyxl import Workbook

        workbook = Workbook()
        for _ in range(8):
            workbook.active.append(["word " * 4000])
        workbook.save(source)
        workbook.close()
    text = extract_context(source)
    assert len(text) == 100_000
    assert "word word" in text


def test_xlsx_formulas_and_links_are_inert_text(tmp_path, monkeypatch):
    from openpyxl import Workbook

    source = tmp_path / "links.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = '=WEBSERVICE("https://example.invalid/private")'
    workbook.active["A2"] = "External link label"
    workbook.active["A2"].hyperlink = "https://example.invalid/private"
    workbook.save(source)
    workbook.close()

    def forbidden(*args, **kwargs):
        pytest.fail("Extraction attempted network access or code execution")

    monkeypatch.setattr(socket.socket, "connect", forbidden)
    monkeypatch.setattr(subprocess, "run", forbidden)
    text = extract_context(source)
    assert '=WEBSERVICE("https://example.invalid/private")' in text
    assert "External link label" in text


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_office_xml_entities_are_rejected_before_external_access(tmp_path, monkeypatch, encoding):
    source = office_file(tmp_path / "entities.docx")
    xml = (
        '<?xml version="1.0" encoding="' + encoding + '"?>'
        '<!DOCTYPE document [<!ENTITY secret SYSTEM "file:///etc/passwd">]>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:body><w:p><w:r><w:t>&secret;</w:t></w:r></w:p></w:body></w:document>'
    )
    rewrite_zip(source, {"word/document.xml": xml.encode(encoding)})
    with pytest.raises(RuntimeError, match="(?i)(XML|entit|DTD).*download-only"):
        extract_context(source)


@pytest.mark.parametrize("limit, value", [
    ("MAX_ZIP_MEMBERS", 10),
    ("MAX_MEMBER_BYTES", 1024),
    ("MAX_EXPANDED_BYTES", 1024),
    ("MAX_INPUT_BYTES", 1024),
    ("MAX_XML_ELEMENTS", 20),
    ("MAX_XML_DEPTH", 3),
])
def test_office_resource_budgets_reject_before_parsing(tmp_path, monkeypatch, limit, value):
    from lecnote import document_extraction

    source = office_file(tmp_path / "bounded.docx")
    monkeypatch.setattr(document_extraction, limit, value)
    with pytest.raises(RuntimeError, match="(?i)(limit|large|many|deep).*download-only"):
        extract_context(source)


def test_highly_compressed_office_member_is_rejected(tmp_path):
    source = office_file(tmp_path / "bomb.docx")
    rewrite_zip(source, additions=[("word/large.xml", b" " * (9 * 1024 * 1024))])
    with pytest.raises(RuntimeError, match="(?i)(limit|large).*download-only"):
        extract_context(source)


def test_encrypted_zip_member_is_rejected(tmp_path):
    source = office_file(tmp_path / "encrypted.docx")
    raw = bytearray(source.read_bytes())
    offset = raw.index(b"PK\x01\x02")
    struct.pack_into("<H", raw, offset + 8, struct.unpack_from("<H", raw, offset + 8)[0] | 1)
    source.write_bytes(raw)
    with pytest.raises(RuntimeError, match="(?i)encrypted.*download-only"):
        extract_context(source)


def test_sparse_spreadsheet_does_not_expand_billions_of_empty_cells(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "sparse.xlsx"
    workbook = Workbook()
    workbook.active["XFD1048576"] = "Far away"
    workbook.save(source)
    workbook.close()
    with pytest.raises(RuntimeError, match="(?i)(limit|large).*download-only"):
        extract_context(source)


def test_pdf_and_ocr_output_are_bounded(tmp_path, monkeypatch):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    from lecnote import enrichment

    source = tmp_path / "large.pdf"
    writer = PdfWriter()
    for _ in range(3):
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
        })
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf (" + b"word " * 10_000 + b") Tj ET")
        page[NameObject("/Contents")] = stream
    with source.open("wb") as output:
        writer.write(output)
    assert len(extract_context(source)) == 100_000
    source = tmp_path / "image.png"
    source.write_bytes(b"image passed to external OCR boundary")
    monkeypatch.setattr(enrichment, "_extract_image", lambda path: "word " * 40_000)
    assert len(extract_context(source)) == 100_000


@pytest.mark.parametrize("attribute", ["gridSpan", "gridBefore", "gridAfter"])
def test_docx_rejects_large_table_dimensions_before_grid_expansion(tmp_path, attribute):
    source = office_file(tmp_path / "table.docx")
    with zipfile.ZipFile(source) as package:
        xml = package.read("word/document.xml")
    if attribute == "gridSpan":
        xml = xml.replace(b"<w:tcPr>", b'<w:tcPr><w:gridSpan w:val="10000"/>', 1)
    else:
        xml = xml.replace(b"<w:tr>", b'<w:tr><w:trPr><w:' + attribute.encode() + b' w:val="10000"/></w:trPr>', 1)
    rewrite_zip(source, {"word/document.xml": xml})
    with pytest.raises(RuntimeError, match="(?i)table.*limit.*download-only"):
        extract_context(source)


def test_docx_merged_cells_are_read_once_without_recursive_vertical_expansion(tmp_path):
    source = office_file(tmp_path / "merged.docx")
    xml = (
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:body><w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w="100"/><w:gridCol w:w="100"/></w:tblGrid>'
        '<w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/><w:vMerge w:val="restart"/></w:tcPr>'
        '<w:p><w:r><w:t>Merged content</w:t></w:r></w:p></w:tc></w:tr>'
        + '<w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/><w:vMerge/></w:tcPr><w:p/></w:tc></w:tr>' * 100
        + '</w:tbl><w:p><w:r><w:t>After table</w:t></w:r></w:p></w:body></w:document>'
    )
    rewrite_zip(source, {"word/document.xml": xml.encode()})
    text = extract_context(source)
    assert text.count("Merged content") == 1
    assert text.endswith("After table")


@pytest.mark.parametrize("suffix", [".pdf", ".png"])
def test_pdf_and_images_reject_oversized_input_before_backend(tmp_path, monkeypatch, suffix):
    from lecnote import enrichment

    source = tmp_path / ("large" + suffix)
    source.write_bytes(b"x" * 1025)
    monkeypatch.setattr(enrichment, "MAX_INPUT_BYTES", 1024, raising=False)
    with pytest.raises(RuntimeError, match="(?i)input limit.*download-only"):
        extract_context(source)
